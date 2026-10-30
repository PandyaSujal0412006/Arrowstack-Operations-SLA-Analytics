"""
==============================================================================
PROJECT 2: OPERATIONS SLA ANALYTICS - EXCEL WORKBOOK GENERATOR
Arrowstack Technologies Data Analytics Internship (Batch 2026)
Candidate: Sujal Pandya | Roll ID: AS-2026-92197
==============================================================================
"""

import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_PATH = os.path.join(BASE_DIR, "data", "cleaned_tickets_data.csv")
OUTPUT_EXCEL_PATH = os.path.join(BASE_DIR, "analysis", "operations_sla_analysis.xlsx")

def create_styled_excel():
    df = pd.read_csv(DATA_PATH)
    resolved_df = df[df["is_resolved"] == 1].copy()
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    # Visual Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    kpi_card_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    kpi_breach_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    
    font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_kpi_label = Font(name="Calibri", size=10, bold=True, color="475569")
    font_kpi_val = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    font_kpi_breach = Font(name="Calibri", size=16, bold=True, color="DC2626")
    font_total = Font(name="Calibri", size=11, bold=True, color="000000")
    font_regular = Font(name="Calibri", size=11, color="000000")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    double_bottom_border = Border(
        top=Side(style='thin', color='000000'),
        bottom=Side(style='double', color='000000')
    )

    # ----------------------------------------------------
    # TAB 1: EXECUTIVE SLA DASHBOARD
    # ----------------------------------------------------
    ws_exec = wb.create_sheet(title="Executive SLA Dashboard")
    ws_exec.views.sheetView[0].showGridLines = True
    
    ws_exec["B2"] = "ARROWSTACK OPERATIONS SLA & INCIDENT PERFORMANCE DASHBOARD"
    ws_exec["B2"].font = font_title
    ws_exec["B3"] = "Executive SLA Analytics Model | Candidate: Sujal Pandya (AS-2026-92197) | Synthetic Dataset"
    ws_exec["B3"].font = font_subtitle
    
    # 6 KPI Cards
    cards = [
        ("TOTAL TICKETS LOGGED", "=COUNTA('Cleaned Tickets Data'!A2:A4201)", "#,##0", "B5", "C6", kpi_card_fill, font_kpi_val),
        ("RESOLVED TICKETS", "=COUNTIF('Cleaned Tickets Data'!D2:D4201, \"Resolved\") + COUNTIF('Cleaned Tickets Data'!D2:D4201, \"Closed\")", "#,##0", "E5", "F6", kpi_card_fill, font_kpi_val),
        ("ACTIVE OPEN BACKLOG", "=C6-F6", "#,##0", "H5", "I6", kpi_card_fill, font_kpi_val),
        ("SLA COMPLIANCE RATE", "=COUNTIF('Cleaned Tickets Data'!L2:L4201, \"Compliant\")/F6", "0.0%", "K5", "L6", kpi_card_fill, font_kpi_val),
        ("SLA BREACH RATE", "=1-L6", "0.0%", "N5", "O6", kpi_breach_fill, font_kpi_breach),
        ("MEAN RESOLUTION TIME (MTTR)", "=AVERAGE('Cleaned Tickets Data'!K2:K4201)", "0.0\" hrs\"", "Q5", "R6", kpi_card_fill, font_kpi_val)
    ]
    
    for title, formula, num_fmt, top_left, bottom_right, fill_type, val_font in cards:
        tl_col, tl_row = top_left[0], int(top_left[1:])
        br_col, br_row = bottom_right[0], int(bottom_right[1:])
        
        ws_exec.merge_cells(f"{top_left}:{chr(ord(tl_col)+1)}{tl_row}")
        c_tl = ws_exec[top_left]
        c_tl.value = title
        c_tl.font = font_kpi_label
        c_tl.alignment = Alignment(horizontal="center", vertical="center")
        
        ws_exec.merge_cells(f"{tl_col}{tl_row+1}:{bottom_right}")
        val_cell = ws_exec[f"{tl_col}{tl_row+1}"]
        val_cell.value = formula
        val_cell.font = val_font
        val_cell.number_format = num_fmt
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for r in range(tl_row, br_row + 1):
            for c in range(ord(tl_col) - ord('A') + 1, ord(br_col) - ord('A') + 1):
                cell = ws_exec.cell(row=r, column=c)
                cell.fill = fill_type
                cell.border = thin_border

    # Priority Summary Table on Executive Sheet
    ws_exec["B9"] = "SLA Compliance Matrix by Ticket Priority"
    ws_exec["B9"].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    
    headers_pri = ["Priority Level", "SLA Target", "Total Tickets", "Resolved", "Compliant", "Breached", "Compliance %", "MTTR (Hours)"]
    for col_idx, h in enumerate(headers_pri, start=2):
        cell = ws_exec.cell(row=10, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
        
    priorities_list = [("P1 - Critical", "4 hrs"), ("P2 - High", "12 hrs"), ("P3 - Medium", "48 hrs"), ("P4 - Low", "96 hrs")]
    for idx, (p_name, target_str) in enumerate(priorities_list, start=11):
        ws_exec.cell(row=idx, column=2, value=p_name).font = font_regular
        ws_exec.cell(row=idx, column=3, value=target_str).alignment = Alignment(horizontal="center")
        ws_exec.cell(row=idx, column=4, value=f"=COUNTIF('Cleaned Tickets Data'!E:E, \"{p_name}\")").number_format = "#,##0"
        ws_exec.cell(row=idx, column=5, value=f"=COUNTIFS('Cleaned Tickets Data'!E:E, \"{p_name}\", 'Cleaned Tickets Data'!O:O, 1)").number_format = "#,##0"
        ws_exec.cell(row=idx, column=6, value=f"=COUNTIFS('Cleaned Tickets Data'!E:E, \"{p_name}\", 'Cleaned Tickets Data'!L:L, \"Compliant\")").number_format = "#,##0"
        ws_exec.cell(row=idx, column=7, value=f"=COUNTIFS('Cleaned Tickets Data'!E:E, \"{p_name}\", 'Cleaned Tickets Data'!L:L, \"Breached\")").number_format = "#,##0"
        ws_exec.cell(row=idx, column=8, value=f"=F{idx}/E{idx}").number_format = "0.0%"
        ws_exec.cell(row=idx, column=9, value=f"=AVERAGEIFS('Cleaned Tickets Data'!K:K, 'Cleaned Tickets Data'!E:E, \"{p_name}\", 'Cleaned Tickets Data'!O:O, 1)").number_format = "0.0"
        
        for col_idx in range(2, 10):
            ws_exec.cell(row=idx, column=col_idx).border = thin_border
            if idx % 2 == 0:
                ws_exec.cell(row=idx, column=col_idx).fill = zebra_fill
                
    # Total Row
    tot_row = 15
    ws_exec.cell(row=tot_row, column=2, value="Grand Total").font = font_total
    ws_exec.cell(row=tot_row, column=3, value="—").alignment = Alignment(horizontal="center")
    ws_exec.cell(row=tot_row, column=4, value="=SUM(D11:D14)").number_format = "#,##0"
    ws_exec.cell(row=tot_row, column=5, value="=SUM(E11:E14)").number_format = "#,##0"
    ws_exec.cell(row=tot_row, column=6, value="=SUM(F11:F14)").number_format = "#,##0"
    ws_exec.cell(row=tot_row, column=7, value="=SUM(G11:G14)").number_format = "#,##0"
    ws_exec.cell(row=tot_row, column=8, value="=F15/E15").number_format = "0.0%"
    ws_exec.cell(row=tot_row, column=9, value="=AVERAGE('Cleaned Tickets Data'!K2:K4201)").number_format = "0.0"
    
    for col_idx in range(2, 10):
        c = ws_exec.cell(row=tot_row, column=col_idx)
        c.font = font_total
        c.fill = total_fill
        c.border = double_bottom_border

    # ----------------------------------------------------
    # TAB 2: TEAM PERFORMANCE
    # ----------------------------------------------------
    ws_team = wb.create_sheet(title="Team Performance")
    ws_team.views.sheetView[0].showGridLines = True
    ws_team["B2"] = "OPERATIONAL TEAM THROUGHPUT & QUEUE PERFORMANCE"
    ws_team["B2"].font = font_title
    
    headers_team = ["Support Team", "Total Tickets", "Resolved", "Open Backlog", "Compliant", "Breached", "Compliance %", "MTTR (Hours)", "Avg CSAT (1-5)"]
    for col_idx, h in enumerate(headers_team, start=2):
        cell = ws_team.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
        
    teams = sorted(df["assigned_team"].unique())
    for r_idx, t in enumerate(teams, start=5):
        ws_team.cell(row=r_idx, column=2, value=t)
        ws_team.cell(row=r_idx, column=3, value=f"=COUNTIF('Cleaned Tickets Data'!G:G, \"{t}\")").number_format = "#,##0"
        ws_team.cell(row=r_idx, column=4, value=f"=COUNTIFS('Cleaned Tickets Data'!G:G, \"{t}\", 'Cleaned Tickets Data'!O:O, 1)").number_format = "#,##0"
        ws_team.cell(row=r_idx, column=5, value=f"=C{r_idx}-D{r_idx}").number_format = "#,##0"
        ws_team.cell(row=r_idx, column=6, value=f"=COUNTIFS('Cleaned Tickets Data'!G:G, \"{t}\", 'Cleaned Tickets Data'!L:L, \"Compliant\")").number_format = "#,##0"
        ws_team.cell(row=r_idx, column=7, value=f"=COUNTIFS('Cleaned Tickets Data'!G:G, \"{t}\", 'Cleaned Tickets Data'!L:L, \"Breached\")").number_format = "#,##0"
        ws_team.cell(row=r_idx, column=8, value=f"=F{r_idx}/D{r_idx}").number_format = "0.0%"
        ws_team.cell(row=r_idx, column=9, value=f"=AVERAGEIFS('Cleaned Tickets Data'!K:K, 'Cleaned Tickets Data'!G:G, \"{t}\", 'Cleaned Tickets Data'!O:O, 1)").number_format = "0.0"
        ws_team.cell(row=r_idx, column=10, value=f"=AVERAGEIFS('Cleaned Tickets Data'!M:M, 'Cleaned Tickets Data'!G:G, \"{t}\", 'Cleaned Tickets Data'!O:O, 1)").number_format = "0.00"
        
        for col_idx in range(2, 11):
            ws_team.cell(row=r_idx, column=col_idx).border = thin_border
            if r_idx % 2 == 1:
                ws_team.cell(row=r_idx, column=col_idx).fill = zebra_fill
                
    # ----------------------------------------------------
    # TAB 3: BREACH ROOT CAUSE ANALYSIS
    # ----------------------------------------------------
    ws_root = wb.create_sheet(title="Breach Root Cause Analysis")
    ws_root.views.sheetView[0].showGridLines = True
    ws_root["B2"] = "SLA BREACH PARETO & ROOT CAUSE TAXONOMY"
    ws_root["B2"].font = font_title
    
    headers_root = ["Breach Root Cause Category", "Breach Count", "Share of Breaches %", "Cumulative Share %", "Operational Remediation Action"]
    for col_idx, h in enumerate(headers_root, start=2):
        cell = ws_root.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
        
    breach_df = resolved_df[resolved_df["sla_status"] == "Breached"].copy()
    root_summary = breach_df["breach_root_cause"].value_counts().reset_index()
    root_summary.columns = ["cause", "count"]
    total_breaches = root_summary["count"].sum()
    
    remediations = {
        "Staffing Deficit / Peak Queue Surge": "Introduce dynamic surge staffing rota and queue routing automation.",
        "Third-Party / Cloud Vendor Dependency": "Enforce vendor SLA escalation protocols and multi-cloud fallback.",
        "Incomplete Customer Diagnostics": "Implement mandatory intake form validation with diagnostics capture.",
        "Complex Code Defect / Regression": "Establish dedicated Tier 3 hotfix sprint bandwidth and rollback CI/CD.",
        "Cross-Team Escalation Latency": "Implement automated 2-hour SLA re-assignment triggers between Tier 1 and SRE."
    }
    
    for r_idx, row in root_summary.iterrows():
        excel_row = 5 + r_idx
        ws_root.cell(row=excel_row, column=2, value=row["cause"])
        ws_root.cell(row=excel_row, column=3, value=row["count"]).number_format = "#,##0"
        ws_root.cell(row=excel_row, column=4, value=f"=C{excel_row}/{total_breaches}").number_format = "0.0%"
        ws_root.cell(row=excel_row, column=5, value=f"=SUM($C$5:C{excel_row})/{total_breaches}").number_format = "0.0%"
        ws_root.cell(row=excel_row, column=6, value=remediations.get(row["cause"], "Process review required."))
        
        for col_idx in range(2, 7):
            ws_root.cell(row=excel_row, column=col_idx).border = thin_border
            if r_idx % 2 == 1:
                ws_root.cell(row=excel_row, column=col_idx).fill = zebra_fill

    # ----------------------------------------------------
    # TAB 4: CLEANED TICKETS DATA (Raw Table)
    # ----------------------------------------------------
    ws_data = wb.create_sheet(title="Cleaned Tickets Data")
    ws_data.views.sheetView[0].showGridLines = True
    
    headers_all = list(df.columns)
    for col_idx, h in enumerate(headers_all, start=1):
        cell = ws_data.cell(row=1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
        
    for r_idx, row in df.iterrows():
        excel_row = 2 + r_idx
        for col_idx, val in enumerate(row, start=1):
            c = ws_data.cell(row=excel_row, column=col_idx, value=val)
            c.font = font_regular
            c.border = thin_border
            if col_idx in [10, 11]: # sla_target_hours, actual_resolution_hours
                c.number_format = "0.0"

    # Autofit Column Widths
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len and not val_str.startswith('='):
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    os.makedirs(os.path.dirname(OUTPUT_EXCEL_PATH), exist_ok=True)
    wb.save(OUTPUT_EXCEL_PATH)
    print(f"[OK] Formatted analytical Excel workbook created at {OUTPUT_EXCEL_PATH}")

if __name__ == "__main__":
    create_styled_excel()
